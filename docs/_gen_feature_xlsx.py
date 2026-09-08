# -*- coding: utf-8 -*-
"""生成按 V1.0.0 / V2.0.0 拆分的功能清单 Excel。"""
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from pathlib import Path
from collections import defaultdict

rows = []  # (ver, m1, m2, feat, desc, fe, be, pri)

# V1：目标/计划/任务相关（含汇报提醒与最小系统壳）
V1_KEYS = {
    ('经营台账', '任务台账'),
    ('经营台账', '工作周报'),
    ('目标计划', '目标管理'),
    ('目标计划', '计划管理'),
    ('任务中心', '任务清单'),
    ('任务中心', '任务督办'),
    ('任务中心', '专项管理'),
    ('任务中心', '汇报提醒'),  # 支撑工作周报催填
    ('系统支撑', 'PC壳'),
    ('系统集成', '组织人员'),
    ('系统集成', '单点登录'),
}


def a(m1, m2, feat, desc, fe, be, pri='P0'):
    ver = 'V1.0.0' if (m1, m2) in V1_KEYS else 'V2.0.0'
    rows.append((ver, m1, m2, feat, desc, fe, be, pri))


# ========== 经营台账 ==========
a('经营台账', '经营总览', '开累底盘指标', '吞吐量/箱量/营收/应收开累看板，不受期间筛选影响', 5, 4)
a('经营台账', '经营总览', '期间作业量看板', '散货/件杂/集装箱KPI、进出口结构、目标环、趋势与货种排行', 10, 8)
a('经营台账', '经营总览', '营收财务看板', '期间营收/开票/回款/应收、货种营收、客户营收排行', 8, 6)
a('经营台账', '经营总览', '应收与超期看板', '账龄分布、客户应收压力、催收Top、超期合同付款条件表', 8, 8)
a('经营台账', '经营总览', '客户/分组明细表', '作业量+开票回款应收超期；支持客户/分组切换与排序', 6, 6)
a('经营台账', '经营总览', '时间筛选联动', '本年/本季/本月共用筛选，驱动作业量与营收区间', 3, 4)
a('经营台账', '经营总览', '明细下钻弹窗', 'KPI/排行点击打开汇总明细弹窗', 4, 3)
a('经营台账', '任务台账', '部门工作台总览', '部门健康分布、待跟进/待计划、部门对比、目标树入口', 8, 6)
a('经营台账', '任务台账', '部门筛选与下钻', '领导切部门；卡片下钻目标/计划/任务', 4, 4)
a('经营台账', '工作周报', '周报/月报汇总页', '目标健康、计划节点、任务落地汇总与较上期对比', 8, 8)
a('经营台账', '工作周报', '填报覆盖与催填入口', '覆盖摘要；链到汇报提醒催填', 3, 4)

# ========== 生产作业 ==========
a('生产作业', '作业量管理', '按日登记作业量', '业务日+货类+货种+客户；散货/件杂万吨；集装箱先TEU后折合万吨', 8, 8)
a('生产作业', '作业量管理', '日明细查询筛选', '日期范围/货类/客户/客户组筛选、编辑删除、唯一键覆盖确认', 5, 5)
a('生产作业', '作业量管理', '月季年汇总', '自然月/季/年聚合；默认客户维，可切客户组；下钻日明细', 8, 10)
a('生产作业', '作业量管理', '供数经营总览', '汇总接口供总览吞吐/排行/客户明细作业量列', 2, 5)
a('生产作业', '货种字典', '货种主数据维护', '货种与货类关系、启停用', 3, 3, 'P1')

# ========== 经营计划 ==========
a('经营计划', '经营目标', '目标值维护', '吞吐/散货/件杂/箱/营收等年季月目标CRUD', 5, 5)
a('经营计划', '经营目标', '完成率对照', '实际取作业量/营收汇总，计算完成率供总览目标环', 3, 4)

# ========== 客户CRM ==========
a('客户CRM', '客户概览', 'CRM经营概览', '客户数、分组分布、跟进活跃度、商机漏斗摘要', 6, 5)
a('客户CRM', '客户管理', '客户档案CRUD', '客户建档、类型等级状态、货种、预估吞吐、负责人', 8, 8)
a('客户CRM', '客户管理', '客户详情360', '基础信息、联系人、洽谈、商机、合同、作业量、应收一页通', 10, 10)
a('客户CRM', '客户管理', '客户导入导出', 'Excel批量导入导出与校验', 5, 6, 'P1')
a('客户CRM', '客户管理', '客户合并去重', '疑似重复识别与合并', 4, 6, 'P1')
a('客户CRM', '分组管理', '客户分组CRUD', '分组创建、成员维护、未分组处理', 4, 4)
a('客户CRM', '分组管理', '分组经营视图', '按组看作业量/营收/应收（对接总览分组维）', 4, 5)
a('客户CRM', '联系人管理', '联系人CRUD', '多联系人、主联系人、职务手机邮箱', 4, 4)
a('客户CRM', '洽谈记录', '写跟进/洽谈', '洽谈新建编辑、关联客户、下次跟进日', 5, 5)
a('客户CRM', '洽谈记录', '跟进提醒列表', '到期未跟进清单与筛选', 3, 4)
a('客户CRM', '商机管理', '商机漏斗', '商机阶段、金额、赢单输单、阶段变更', 10, 10)
a('客户CRM', '商机管理', '商机转合同', '赢单商机一键生成合同草稿', 3, 4)
a('客户CRM', '公海私海', '客户池规则', '领取/回收/掉入公海规则', 6, 8, 'P1')
a('客户CRM', '客户拜访', '拜访计划与记录', '计划拜访、签到记录', 5, 5, 'P2')

# ========== 合同管理 ==========
a('合同管理', '合同台账', '合同全生命周期', '起草/执行中/变更/完成/终止/作废状态机与列表', 10, 12)
a('合同管理', '合同台账', '合同新建编辑', '名称编号类型客户金额期间责任部门说明', 6, 6)
a('合同管理', '合同台账', '付款条件与账期', '付款条件、应付日起算规则；未约定不进超期表', 5, 6)
a('合同管理', '合同台账', '计费明细', '铁水联运等费项多行（装卸/堆存/代理）', 5, 5)
a('合同管理', '合同台账', '合同详情总览', '金额/开票/回款/剩余、账期、关联开票收款、附件入口', 6, 6)
a('合同管理', '合同台账', '合同变更', '金额/账期/期限变更记录与版本留痕', 6, 8)
a('合同管理', '合同台账', '合同提醒', '到期前提醒、账期临近提醒', 3, 4, 'P1')
a('合同管理', '相对方', '相对方档案关联', '合同相对方与客户主数据绑定校验', 3, 3)
a('合同管理', '合同模板', '模板与条款库', '常用合同模板套用', 5, 5, 'P1')
a('合同管理', '电子签', '签署对接预留', '对接电子签平台接口预留', 3, 8, 'P2')

# ========== 结算收款 ==========
a('结算收款', '开票管理', '开票登记', '必须关联合同；票号类型金额日期；超额开票确认', 6, 6)
a('结算收款', '开票管理', '发票作废', '仅改状态已作废，不计入累计开票', 3, 3)
a('结算收款', '开票管理', '开票查询统计', '按年/合同/状态筛选与累计待开票', 4, 4)
a('结算收款', '收款管理', '收款登记', '必须关联合同；可选关联开票；无强关联', 6, 6)
a('结算收款', '收款管理', '收款查询统计', '回款清单、超额回款提示、计划收款', 4, 5)
a('结算收款', '应收超期', '应收与超期计算', '应收=营收-回款或合同未收（可配置）；相对应付日算超期', 4, 8)
a('结算收款', '应收超期', '催收清单', '超期客户/合同催收列表与导出', 4, 4)
a('结算收款', '营收台账', '经营营收登记', '按客户/货类期间营收登记（经营口径，供总览）', 6, 6)

# ========== 目标计划 ==========
a('目标计划', '目标管理', '目标树列表', '年-季-月树形；健康状态；量化达成；拆解下级', 10, 10)
a('目标计划', '目标管理', '目标新建编辑详情', '介绍、截止、责任部门、量化指标、改状态必填原因', 8, 8)
a('目标计划', '计划管理', '计划列表与CRUD', '关联目标、关键节点、进度（节点为主任务为辅）', 8, 8)
a('目标计划', '计划管理', '计划详情与节点', '节点达成、下属任务概况', 5, 5)

# ========== 任务中心 ==========
a('任务中心', '任务清单', '我的任务列表', '本人任务筛选、状态、滞后标识', 6, 6)
a('任务中心', '任务清单', '任务详情大弹窗', '基础信息、进展填报、附件、评论编辑删除', 10, 10)
a('任务中心', '任务清单', '任务新建编辑', '关联计划、周期、图片、责任人', 6, 6)
a('任务中心', '任务督办', '督办下达', '领导下达督办任务给指定人员', 5, 5)
a('任务中心', '汇报提醒', '未填报清单', '周/月未填人员；一键催填（领导/部门负责人）', 5, 6)
a('任务中心', '专项管理', '专项列表与详情', '专项状态页签、关联计划任务、删除级联确认', 8, 8)

# ========== 档案 / 知识库 / AI ==========
a('档案管理', '档案库', '档案目录与分类', '按业务域/年度/密级建立目录树', 6, 6)
a('档案管理', '档案库', '档案条目CRUD', '题名、文号、归档日期、保管期限、关联业务对象', 8, 8)
a('档案管理', '档案库', '附件上传预览', '文档/图片上传、在线预览、下载、版本', 8, 8)
a('档案管理', '档案库', '关联挂接', '挂接客户/合同/专项/任务等业务单据', 4, 6)
a('档案管理', '档案库', '全文检索', '题名/文号/正文OCR或文本检索', 4, 8, 'P1')
a('档案管理', '借阅管理', '借阅申请与归还', '借阅流程记录、到期催还（审批走平台）', 5, 6, 'P1')
a('档案管理', '归档规则', '业务自动归档', '合同生效/完结等事件自动归档规则', 3, 6, 'P1')
a('档案管理', '权限与水印', '档案权限与水印', '按密级/部门可见；预览水印', 4, 5, 'P1')
a('知识库', '知识空间', '空间与分类', '知识库空间、栏目、标签', 5, 5)
a('知识库', '文档管理', '文档发布', '富文本/Markdown发布、草稿、上下架', 8, 6)
a('知识库', '文档管理', '文档协作', '评论、点赞收藏、相关推荐', 4, 4, 'P1')
a('知识库', '文档管理', '附件与版本', '附件、历史版本对比回滚', 5, 6)
a('知识库', '检索', '知识检索', '关键词/标签检索、高亮', 4, 5)
a('知识库', '权限', '可见范围', '全员/部门/角色可见配置（对接平台）', 3, 4)
a('知识库', '运营', '阅读统计', '阅读量、热门文档', 2, 3, 'P1')
a('AI智能助手', '对话助手', '统一对话入口', '侧栏/悬浮窗多轮对话；会话历史', 8, 8)
a('AI智能助手', '对话助手', '经营问答', '问吞吐/应收/超期/客户等，基于权限内数据回答', 6, 12)
a('AI智能助手', '对话助手', '引用溯源', '回答附数据来源链接（总览/合同/客户）', 4, 6)
a('AI智能助手', '知识问答', '知识库RAG', '基于知识库/档案检索增强生成', 8, 15)
a('AI智能助手', '业务助理', '填报辅助', '辅助生成周报摘要、任务进展草稿', 5, 8, 'P1')
a('AI智能助手', '业务助理', '催收建议', '基于超期清单生成催收优先级建议', 4, 6, 'P1')
a('AI智能助手', '安全合规', '权限与脱敏', '按数据权限过滤检索；敏感字段脱敏', 4, 8)
a('AI智能助手', '运营配置', '提示词与知识源配置', '可配业务提示词、挂接知识源与表权限', 5, 6, 'P1')
a('AI智能助手', '评测', '问答评测集', '关键问法回归与人工评分', 3, 4, 'P2')

# ========== 系统 ==========
a('系统集成', '组织人员', '组织同步', '同步平台组织/部门/人员', 2, 6)
a('系统集成', '单点登录', 'SSO登录', '对接内部平台登录鉴权', 2, 5)
a('系统集成', '外部数据', '生产/财务同步预留', '作业量、开票回款外部系统同步适配', 3, 10, 'P1')
a('系统支撑', '基础设置', '字典与枚举', '货类货种、合同类型、付款条件等字典', 4, 4)
a('系统支撑', '操作日志', '关键操作审计', '合同/收款/档案等关键操作日志', 3, 4)
a('系统支撑', '消息对接', '待办与通知', '催填/到期/超期通过平台消息下发', 2, 5, 'P1')
a('系统支撑', 'PC壳', '统一布局菜单', '侧栏菜单、顶栏、通用弹窗全屏', 4, 2)
a('系统支撑', '企微H5', '企微关键页适配', '总览/任务台账/周报/我的任务等H5', 10, 6, 'P1')
a('经营分析', '报表中心', '标准经营报表', '月度吞吐/营收/回款/超期报表导出', 6, 8, 'P1')
a('经营分析', '报表中心', '自定义查询', '多维筛选导出', 5, 8, 'P2')

# ---------- Excel ----------
header_fill = PatternFill('solid', fgColor='0F766E')
header_font = Font(color='FFFFFF', bold=True, size=11, name='微软雅黑')
pri_fills = {
    'P0': PatternFill('solid', fgColor='FEF3C7'),
    'P1': PatternFill('solid', fgColor='E0F2FE'),
    'P2': PatternFill('solid', fgColor='F1F5F9'),
}
ver_fills = {
    'V1.0.0': PatternFill('solid', fgColor='D1FAE5'),
    'V2.0.0': PatternFill('solid', fgColor='E0E7FF'),
}
thin = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1'),
)
wrap = Alignment(wrap_text=True, vertical='center')

HEADERS = [
    '序号', '版本', '一级模块', '二级模块', '功能点', '说明', '优先级',
    '前端人天', '后端人天', '预估工时(人天)',
    '2前端并行日历天', '2后端并行日历天', '建议排期参考(周)'
]


def write_feature_sheet(ws, data, start_seq=1):
    for col, h in enumerate(HEADERS, 1):
        cell = ws.cell(1, col, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for i, (ver, m1, m2, feat, desc, fe, be, pri) in enumerate(data, start_seq):
        total = fe + be
        fe_cal = round(fe / 2, 1)
        be_cal = round(be / 2, 1)
        weeks = round(max(fe_cal, be_cal) / 5, 1)
        vals = [i, ver, m1, m2, feat, desc, pri, fe, be, total, fe_cal, be_cal, weeks]
        row_idx = i - start_seq + 2
        for col, v in enumerate(vals, 1):
            cell = ws.cell(row_idx, col, v)
            cell.border = thin
            cell.alignment = wrap
            cell.font = Font(name='微软雅黑', size=10)
            if col == 2:
                cell.fill = ver_fills.get(ver, ver_fills['V2.0.0'])
                cell.alignment = Alignment(horizontal='center', vertical='center')
            if col == 7:
                cell.fill = pri_fills.get(pri, pri_fills['P2'])
                cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 36
    widths = [6, 10, 14, 14, 22, 46, 8, 10, 10, 12, 14, 14, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = f'A1:M{len(data) + 1}'
    ws.freeze_panes = 'A2'


def write_summary_sheet(ws, data, title_prefix=''):
    headers = [
        '一级模块', '功能点数', 'P0', 'P1', 'P2',
        '前端人天', '后端人天', '预估总人天',
        '2前端日历天', '2后端日历天', '关键路径约周'
    ]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    agg = defaultdict(lambda: {'n': 0, 'p0': 0, 'p1': 0, 'p2': 0, 'fe': 0, 'be': 0})
    for ver, m1, m2, feat, desc, fe, be, pri in data:
        agg[m1]['n'] += 1
        agg[m1][pri.lower()] += 1
        agg[m1]['fe'] += fe
        agg[m1]['be'] += be
    r = 2
    fe_all = be_all = 0
    for m1, d in agg.items():
        fe_all += d['fe']
        be_all += d['be']
        fe_cal = round(d['fe'] / 2, 1)
        be_cal = round(d['be'] / 2, 1)
        vals = [
            m1, d['n'], d['p0'], d['p1'], d['p2'], d['fe'], d['be'], d['fe'] + d['be'],
            fe_cal, be_cal, round(max(fe_cal, be_cal) / 5, 1)
        ]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            cell.border = thin
            cell.font = Font(name='微软雅黑', size=10)
            cell.alignment = wrap
        r += 1
    vals = [
        '合计', len(data),
        sum(1 for x in data if x[7] == 'P0'),
        sum(1 for x in data if x[7] == 'P1'),
        sum(1 for x in data if x[7] == 'P2'),
        fe_all, be_all, fe_all + be_all,
        round(fe_all / 2, 1), round(be_all / 2, 1),
        round(max(fe_all / 2, be_all / 2) / 5, 1)
    ]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(r, c, v)
        cell.border = thin
        cell.font = Font(name='微软雅黑', size=10, bold=True)
        cell.fill = PatternFill('solid', fgColor='ECFDF5')
    for i, w in enumerate([16, 10, 8, 8, 8, 10, 10, 12, 12, 12, 14], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = 'A2'
    return fe_all, be_all


v1 = [r for r in rows if r[0] == 'V1.0.0']
v2 = [r for r in rows if r[0] == 'V2.0.0']

wb = openpyxl.Workbook()

# 版本对照
ws0 = wb.active
ws0.title = '版本对照'
ws0['A1'] = '经营管理系统 · 版本范围对照'
ws0['A1'].font = Font(name='微软雅黑', size=14, bold=True)
ws0['A3'] = '版本'
ws0['B3'] = '范围说明'
ws0['C3'] = '包含二级模块'
ws0['D3'] = '功能点数'
ws0['E3'] = '预估人天'
for c in range(1, 6):
    ws0.cell(3, c).fill = header_fill
    ws0.cell(3, c).font = header_font
v1_fe = sum(x[5] for x in v1)
v1_be = sum(x[6] for x in v1)
v2_fe = sum(x[5] for x in v2)
v2_be = sum(x[6] for x in v2)
ws0['A4'] = 'V1.0.0'
ws0['B4'] = '第一期：工作目标、计划、任务执行闭环'
ws0['C4'] = '任务台账、工作周报、目标管理、计划管理、任务清单、任务督办、专项管理、汇报提醒；+PC壳/组织/SSO'
ws0['D4'] = len(v1)
ws0['E4'] = v1_fe + v1_be
ws0['A5'] = 'V2.0.0'
ws0['B5'] = '第二期：经营总览与生产/客户/合同/结算/档案/知识库/AI等'
ws0['C5'] = '经营总览、生产作业、经营计划(目标值)、客户CRM、合同、结算、档案、知识库、AI、报表等'
ws0['D5'] = len(v2)
ws0['E5'] = v2_fe + v2_be
ws0['A4'].fill = ver_fills['V1.0.0']
ws0['A5'].fill = ver_fills['V2.0.0']
for col, w in enumerate([12, 42, 70, 10, 12], 1):
    ws0.column_dimensions[get_column_letter(col)].width = w
for r in range(4, 6):
    for c in range(1, 6):
        ws0.cell(r, c).border = thin
        ws0.cell(r, c).alignment = wrap
        ws0.cell(r, c).font = Font(name='微软雅黑', size=10)

ws_all = wb.create_sheet('功能清单-全量')
write_feature_sheet(ws_all, rows)

ws_v1 = wb.create_sheet('功能清单-V1.0.0')
write_feature_sheet(ws_v1, v1)

ws_v2 = wb.create_sheet('功能清单-V2.0.0')
write_feature_sheet(ws_v2, v2)

ws_sum = wb.create_sheet('模块汇总-全量')
write_summary_sheet(ws_sum, rows)
ws_sum1 = wb.create_sheet('模块汇总-V1.0.0')
write_summary_sheet(ws_sum1, v1)
ws_sum2 = wb.create_sheet('模块汇总-V2.0.0')
write_summary_sheet(ws_sum2, v2)

ws3 = wb.create_sheet('编制说明')
notes = [
    '经营管理系统 · 完整功能清单（已按 V1.0.0 / V2.0.0 拆分）',
    '',
    '一、版本策略（2026-09-08 确认）',
    'V1.0.0 第一期：与工作目标、计划、任务相关 —— 任务台账、工作周报、目标管理、计划管理、任务清单、任务督办、专项管理；',
    '另含汇报提醒（支撑周报催填），以及可上线所需最小系统能力（PC壳、组织同步、SSO）。',
    'V2.0.0 第二期：其余全部 —— 经营总览、生产作业、经营目标值、完整CRM、完整合同、结算收款、档案、知识库、AI助手、报表等。',
    '',
    '二、工时口径',
    '前端人天/后端人天为开发工作量（含自测与基本联调）；预估工时=前+后；2人并行日历天=人天÷2；建议周≈max日历天÷5。',
    '不含专职测试、实施培训、大模型采购；消息/审批/角色菜单权限由内部平台提供。建议另加15%～25%联调缓冲。',
    '',
    '三、优先级',
    'P0首期必做；P1增强；P2远期。版本字段与优先级独立：某功能可在V2且为P0。',
]
for i, line in enumerate(notes, 1):
    cell = ws3.cell(i, 1, line)
    bold = i == 1 or line.startswith(('一、', '二、', '三、'))
    cell.font = Font(name='微软雅黑', size=11, bold=bold)
ws3.column_dimensions['A'].width = 120

docs = Path(__file__).resolve().parent
out_en = docs / 'BMS_Feature_List_Full.xlsx'
wb.save(out_en)
for p in docs.glob('*.xlsx'):
    if p.name != out_en.name:
        try:
            p.unlink()
        except OSError:
            pass
desktop = Path.home() / 'Desktop' / 'BMS_Feature_List_Full.xlsx'
try:
    wb.save(desktop)
    desktop_msg = str(desktop)
except PermissionError:
    desktop_msg = 'SKIPPED_LOCKED'

stat = docs / '_feature_stat.txt'
stat.write_text(
    f'ALL count={len(rows)} FE={sum(x[5] for x in rows)} BE={sum(x[6] for x in rows)} TOTAL={sum(x[5]+x[6] for x in rows)}\n'
    f'V1 count={len(v1)} FE={v1_fe} BE={v1_be} TOTAL={v1_fe+v1_be} weeks={round(max(v1_fe,v1_be)/2/5,1)}\n'
    f'V2 count={len(v2)} FE={v2_fe} BE={v2_be} TOTAL={v2_fe+v2_be} weeks={round(max(v2_fe,v2_be)/2/5,1)}\n'
    f'out={out_en}\ndesktop={desktop_msg}\n',
    encoding='utf-8'
)
